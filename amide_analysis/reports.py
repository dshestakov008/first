"""Lossless numeric tables and formula-safe, auditable CSV/Excel exports."""

from __future__ import annotations

import json
import re
from itertools import chain
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


MAIN_SHEETS = (
    "Metadata", "Raw_Spectra", "Processed_Spectra", "Second_Derivative",
    "Peak_Parameters", "Secondary_Structure", "Group_Statistics",
    "Statistical_Tests", "Fit_Quality", "Processing_Settings",
)
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
EXCEL_MAX_TEXT_LENGTH = 32_767


def _json_default(value: Any):
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, np.ndarray):
        return value.tolist()
    raise TypeError(f"Cannot serialise setting of type {type(value).__name__}")


def _safe_text(value: Any):
    """Protect text formulas; real numeric values stay numeric and unrounded."""
    if isinstance(value, (dict, list, tuple, np.ndarray)):
        value = json.dumps(value, ensure_ascii=False, default=_json_default)
    if isinstance(value, pd.Timestamp) and value.tzinfo is not None:
        value = value.isoformat()
    if not isinstance(value, str):
        return value
    # Excel XML cannot contain these controls. Explicit escapes retain evidence.
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]",
                   lambda match: f"\\u{ord(match.group()):04x}", value)
    if value.lstrip().lstrip("\ufeff").lstrip().startswith(("=", "+", "-", "@")):
        value = "'" + value
    return value


def _safe_frame(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    for column in result.columns:
        if isinstance(result[column].dtype, pd.CategoricalDtype):
            result[column] = result[column].astype(object).map(_safe_text)
        elif pd.api.types.is_object_dtype(result[column].dtype) or pd.api.types.is_string_dtype(result[column].dtype):
            result[column] = result[column].map(_safe_text)
    result.columns = [_safe_text(str(column)) for column in result.columns]
    return result


def _settings(config: dict) -> pd.DataFrame:
    rows = [("configuration_json", json.dumps(config, ensure_ascii=False, sort_keys=True,
                                             default=_json_default))]

    def visit(value, prefix=""):
        if isinstance(value, dict):
            for key, child in value.items():
                visit(child, f"{prefix}.{key}" if prefix else str(key))
        else:
            encoded = (json.dumps(value, ensure_ascii=False, default=_json_default)
                       if isinstance(value, (list, tuple)) else value)
            rows.append((prefix, encoded))

    visit(config)
    rows.append(("export_text_safety", "Text starting with =, +, -, or @ is prefixed with an apostrophe; numeric cells remain numeric."))
    rows.append(("export_numeric_precision", "Values are exported without rounding. CSV retains full float precision; Excel has its standard numeric precision limits."))
    return pd.DataFrame(rows, columns=["key", "value"])


def _csv_name(name: str) -> str:
    name = re.sub(r"[\\/\x00-\x1f<>:*?\"|]", "_", name).strip(" .")
    return name or "Table"


def _sheet_name(name: str, used: set[str]) -> str:
    stem = re.sub(r"[\\/*?:\[\]]", "_", name).strip("'") or "Table"
    candidate = stem[:31]
    index = 2
    while candidate.casefold() in used:
        suffix = f"_{index}"
        candidate = stem[:31 - len(suffix)] + suffix
        index += 1
    used.add(candidate.casefold())
    return candidate


def export_results(output_dir: Path, tables: dict[str, pd.DataFrame], config: dict) -> Path:
    """Save ten specified sheets, optional diagnostic tables, and full CSVs.

    Every principal table receives its own CSV, including a table with no rows.
    Oversized tables are fully retained as CSV and raise an explicit exception:
    Excel is never silently truncated or represented by a misleading subset.
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    ordered = {name: tables.get(name, pd.DataFrame()) for name in MAIN_SHEETS}
    ordered.update({name: frame for name, frame in tables.items() if name not in ordered})
    for name, frame in ordered.items():
        if not isinstance(frame, pd.DataFrame):
            raise TypeError(f"Table {name!r} must be a pandas DataFrame")
        if not frame.columns.is_unique:
            raise ValueError(f"Table {name!r} contains duplicate column names")
    settings = _settings(config)
    provided = ordered["Processing_Settings"]
    if not provided.empty:
        if {"key", "value"}.issubset(provided.columns):
            existing = set(provided["key"].astype(str))
            settings = pd.concat([provided, settings.loc[~settings["key"].isin(existing)]], ignore_index=True)
        else:
            # Preserve supplied columns and add auditable config rows explicitly.
            settings = pd.concat([provided, settings], ignore_index=True)
    ordered["Processing_Settings"] = settings
    safe = {name: _safe_frame(frame) for name, frame in ordered.items()}
    csv_names: set[str] = set()
    for name, frame in safe.items():
        stem = _csv_name(name)
        candidate, index = stem, 2
        while candidate.casefold() in csv_names:
            candidate = f"{stem}_{index}"
            index += 1
        csv_names.add(candidate.casefold())
        frame.to_csv(output_dir / f"{candidate}.csv", index=False, encoding="utf-8-sig")
    for name, frame in safe.items():
        if len(frame) + 1 > EXCEL_MAX_ROWS:
            raise ValueError(f"{name}: {len(frame)} rows exceed Excel's {EXCEL_MAX_ROWS - 1} data-row limit. Full CSV files were saved; the workbook was not truncated.")
        if len(frame.columns) > EXCEL_MAX_COLUMNS:
            raise ValueError(f"{name}: too many columns for Excel. Full CSV files were saved; the workbook was not truncated.")
        text_values = (value for column in frame.select_dtypes(include=["object", "string"])
                       for value in frame[column])
        if any(isinstance(value, str) and len(value) > EXCEL_MAX_TEXT_LENGTH
               for value in chain(frame.columns, text_values)):
            raise ValueError(f"{name}: a text cell exceeds Excel's {EXCEL_MAX_TEXT_LENGTH}-character limit. Full CSV files were saved; the workbook was not truncated.")
    destination = output_dir / "Amide_I_Results.xlsx"
    temporary = output_dir / ".Amide_I_Results.writing.xlsx"
    used: set[str] = set()
    try:
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            for name, frame in safe.items():
                sheet = _sheet_name(name, used)
                frame.to_excel(writer, sheet_name=sheet, index=False)
                worksheet = writer.sheets[sheet]
                worksheet.freeze_panes = "A2"
                if len(frame.columns):
                    worksheet.auto_filter.ref = worksheet.dimensions
                for cell in worksheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="30485E")
                for cell in worksheet[1]:
                    header = str(cell.value or "")
                    worksheet.column_dimensions[cell.column_letter].width = min(45, max(14, len(header) + 2))
                if name == "Region_Summary":
                    widths = (62, 18, 18, 18, 16, 16, 12)
                    for index, width in enumerate(widths, start=1):
                        worksheet.column_dimensions[worksheet.cell(1, index).column_letter].width = width
                    worksheet.row_dimensions[1].height = 32
                    for cell in worksheet[1]:
                        cell.font = Font(name="Times New Roman", size=11, bold=True, color="222222")
                        cell.fill = PatternFill(fill_type=None)
                        cell.border = Border(bottom=Side(style="thin", color="666666"))
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    for row in worksheet.iter_rows(min_row=2):
                        for cell in row:
                            cell.font = Font(name="Times New Roman", size=11)
                            cell.alignment = Alignment(horizontal="left" if cell.column == 1 else "center", vertical="center", wrap_text=True)
                            if cell.column in (3, 4):
                                cell.number_format = "0.0000"
                            elif cell.column in (5, 6):
                                cell.number_format = '[<0.0001]"<0.0001";0.0000'
                            elif cell.column == 7:
                                cell.number_format = "0.000"
                        worksheet.row_dimensions[row[0].row].height = 32
                if name == "Analysis_Summary":
                    for column, width in zip("ABCD", (34, 40, 26, 100)):
                        worksheet.column_dimensions[column].width = width
                    for row in worksheet.iter_rows(min_row=2):
                        worksheet.row_dimensions[row[0].row].height = 44
                        for cell in row:
                            cell.alignment = Alignment(vertical="center", wrap_text=True)
                        target = row[2]
                        if target.value in writer.sheets:
                            target.hyperlink = f"#'{target.value}'!A1"
                            target.style = "Hyperlink"
            if "Analysis_Summary" in writer.sheets:
                writer.book.active = writer.book.sheetnames.index("Analysis_Summary")
        temporary.replace(destination)
    finally:
        if temporary.exists():
            temporary.unlink()
    return destination
