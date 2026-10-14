import json
import re
import sys
import traceback
import math
from itertools import combinations, islice
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, simpledialog, messagebox

try:
    import numpy as np
    import pandas as pd
    import matplotlib
    import matplotlib.pyplot as plt
    import openpyxl
    import scipy
    from scipy import stats
    from scipy.integrate import trapezoid
    from scipy.stats import false_discovery_control
except ImportError as exc:
    print("Не хватает библиотеки:", exc.name)
    print("В командной строке выполните:")
    print('py -m pip install numpy pandas matplotlib openpyxl xlrd "scipy>=1.11"')
    sys.exit(1)


# Это защитные ограничения программы, а не стандарты ИК-спектроскопии.
STEP_TOLERANCE = 0.01       # Неравномерность шага внутри файла: до 1%.
STEP_RATIO_LIMIT = 1.10     # Шаги разных файлов отличаются не более чем на 10%.
METHOD = "L2_without_centering"
FDR_LEVEL = 0.05
MAX_EXACT_PERMUTATIONS = 50000
RANDOM_PERMUTATIONS = 49999
RANDOM_SEED = 20261004
# Каждая область — отдельная заранее заданная гипотеза. Границы в см⁻¹.
# Область 1800–900 намеренно перекрывается с несколькими более узкими областями.
SPECTRAL_REGIONS = (
    ("Амид I", 1600.0, 1700.0),
    ("Амид II", 1500.0, 1600.0),
    ("Амид III и фосфатные группы", 1200.0, 1350.0),
    ("Липиды, белки и карбоксильные группы", 1350.0, 1500.0),
    ("Углеводы, гликопротеины и фосфаты", 900.0, 1200.0),
    ("Карбонильные группы липидов", 1700.0, 1750.0),
    ("C–H-колебания (липиды и органические компоненты)", 2800.0, 3000.0),
    ("O–H и N–H-колебания (вода, белки)", 3000.0, 3600.0),
    ("Область молекулярных отпечатков", 900.0, 1800.0),
)
Y_LABELS = {
    "I": "Normalized intensity (a.u.)",
    "A": "Normalized absorbance (a.u.)",
    "T": "Normalized transmittance (a.u.)",
}


def is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def number(value):
    # Поддерживаются числа Excel и текстовые числа с десятичной запятой.
    if is_blank(value) or isinstance(value, (bool, np.bool_)):
        raise ValueError("пустая ячейка или логическое значение")
    if isinstance(value, str):
        value = value.strip().replace("\u00a0", "").replace("\u202f", "")
        value = value.replace(" ", "").replace("−", "-").replace(",", ".")
    result = float(value)
    if not np.isfinite(result):
        raise ValueError("NaN или бесконечность")
    return result


def is_x_header(value):
    text = str(value).strip().lower()
    return text == "x" or any(word in text for word in
        ("wavenumber", "wave number", "волнов", "cm-1", "cm⁻¹", "см-1", "см⁻¹"))


def is_y_header(value):
    text = str(value).strip().lower()
    return text == "y" or any(word in text for word in
        ("intensity", "absorbance", "transmittance", "signal",
         "интенсив", "поглощ", "пропуск", "сигнал"))


def read_spectrum(path, say=print):
    path = Path(path)
    try:
        if path.name.startswith("~$"):
            raise ValueError("это временный файл Excel. Выберите исходный файл")
        engine = {".xlsx": "openpyxl", ".xls": "xlrd"}.get(path.suffix.lower())
        if engine is None:
            raise ValueError("поддерживаются только .xlsx и .xls")
        # header=None сохраняет первую точку в файлах без заголовков.
        with pd.ExcelFile(path, engine=engine) as book:
            sheets = []
            for sheet_name in book.sheet_names:
                table = pd.read_excel(book, sheet_name=sheet_name, header=None,
                                      dtype=object, na_filter=False)
                occupied = table.apply(lambda col: col.map(lambda v: not is_blank(v)))
                if occupied.to_numpy().any():
                    columns = list(table.columns[occupied.any(axis=0)])
                    rows = list(table.index[occupied.any(axis=1)])
                    sheets.append((sheet_name, table, columns, rows))
        if len(sheets) != 1:
            raise ValueError("нужен один непустой лист с одним спектром")
        sheet_name, table, columns, rows = sheets[0]
        if len(columns) != 2:
            raise ValueError("нужны ровно два заполненных столбца: волновое число и сигнал")
        first, last = rows[0], rows[-1]
        x_col, y_col = columns
        a, b = table.loc[first, [x_col, y_col]]
        header = False
        if is_x_header(a) and is_y_header(b):
            header = True
        elif is_y_header(a) and is_x_header(b):
            x_col, y_col = y_col, x_col
            header = True
        if header:
            first += 1
        x_values, y_values = [], []
        for row in range(first, last + 1):
            try:
                x_values.append(number(table.at[row, x_col]))
                y_values.append(number(table.at[row, y_col]))
            except (ValueError, TypeError, OverflowError) as exc:
                raise ValueError(f"лист '{sheet_name}', строка Excel {row + 1}: "
                                 f"пропуск или нечисловое значение ({exc})") from exc
        x, y = np.array(x_values), np.array(y_values)
        if len(x) < 3:
            raise ValueError("найдено менее трёх точек")
        if np.any(x <= 0):
            raise ValueError("волновые числа должны быть положительными")
        order = np.argsort(x)
        sorted_rows = not np.array_equal(order, np.arange(len(x)))
        x, y = x[order], y[order]
        delta = np.diff(x)
        if np.any(delta == 0):
            raise ValueError("есть повторяющиеся волновые числа. Исправьте исходный файл")
        step = float(np.median(delta))
        if np.max(np.abs(delta - step)) / step > STEP_TOLERANCE:
            raise ValueError(f"шаг неравномерен или есть потерянные точки: "
                             f"от {delta.min():.6g} до {delta.max():.6g} см-1")
        if np.ptp(y) == 0:
            raise ValueError("сигнал одинаков во всех точках. Проверьте экспорт спектра")
        if sorted_rows:
            say("  Точки упорядочены по возрастанию волнового числа.")
        if np.any(y < 0):
            say("  Есть отрицательные значения. Они сохранены без изменений.")
        say(f"  Точек: {len(x)}; диапазон: {x[0]:.6f}–{x[-1]:.6f} см-1")
        say(f"  Шаг: медиана {step:.6f}, минимум {delta.min():.6f}, "
            f"максимум {delta.max():.6f} см-1")
        return {"path": path, "sheet": sheet_name, "x": x, "y": y,
                "step": step, "sorted": sorted_rows, "header": header}
    except Exception as exc:
        raise ValueError(f"Файл '{path.name}': {exc}") from exc


def choose_grid(spectra, reference_path, signal_type, confirm, say):
    # Общая сетка сохраняется и используется при отдельных запусках для других групп.
    is_new = not reference_path.exists()
    if is_new:
        base = max(spectra, key=lambda item: item["step"])
        low = max(item["x"][0] for item in spectra)
        high = min(item["x"][-1] for item in spectra)
        grid = base["x"][(base["x"] >= low) & (base["x"] <= high)].copy()
        say("Создаётся общая сетка на основе файла с наибольшим шагом.")
    else:
        try:
            saved = json.loads(reference_path.read_text(encoding="utf-8"))
            if saved["method"] != METHOD or saved["signal_type"] != signal_type:
                raise ValueError("метод или выбранный тип сигнала отличается от первого запуска")
            grid = np.asarray(saved["wavenumber_cm-1"], dtype=float)
        except Exception as exc:
            raise ValueError(f"Файл общей сетки: {exc}. Для нового эксперимента "
                             "используйте отдельную папку с программой") from exc
        say("Используется общая сетка предыдущего запуска.")
    if (grid.ndim != 1 or len(grid) < 3 or not np.isfinite(grid).all()
            or np.any(np.diff(grid) <= 0)):
        raise ValueError("нет корректной общей сетки: проверьте диапазоны исходных файлов")
    grid_step = float(np.median(np.diff(grid)))
    if np.max(np.abs(np.diff(grid) - grid_step)) / grid_step > STEP_TOLERANCE:
        raise ValueError("общая сетка имеет неравномерный шаг")
    trimmed = False
    for item in spectra:
        x, step = item["x"], item["step"]
        if max(step, grid_step) / min(step, grid_step) > STEP_RATIO_LIMIT:
            raise ValueError(f"Файл '{item['path'].name}': шаг слишком отличается "
                             "от общей сетки. Проверьте параметры экспорта")
        if max(abs(x[0] - grid[0]), abs(x[-1] - grid[-1])) > 2 * max(step, grid_step):
            raise ValueError(f"Файл '{item['path'].name}': диапазон заметно отличается "
                             "от общей сетки. Автоматическое обрезание не выполнено")
        if x[0] > grid[0] or x[-1] < grid[-1]:
            raise ValueError(f"Файл '{item['path'].name}': не покрывает общую сетку "
                             f"{grid[0]:.6f}–{grid[-1]:.6f} см-1. "
                             "Экстраполяция запрещена. Нужен экспорт общего диапазона")
        trimmed |= bool(x[0] < grid[0] or x[-1] > grid[-1])
    if trimmed:
        text = (f"Крайние точки файлов немного различаются.\n"
                f"Общая область: {grid[0]:.6f}–{grid[-1]:.6f} см-1.\n"
                "Некоторые крайние точки вне этой области не войдут в результат.\n"
                "Внутренние области сохранятся полностью. Экстраполяции не будет.\n\n"
                "Разрешить такое ограничение краёв для всей группы?")
        if not confirm(text):
            raise ValueError("ограничение краёв отменено. Обработка остановлена")
        say("Пользователь разрешил ограничить только края общей областью.")
    say(f"Общая сетка: {len(grid)} точек; {grid[0]:.6f}–{grid[-1]:.6f} см-1")
    return grid, is_new


def format_excel(writer):
    # Простое оформление таблиц, без изменения чисел и без округления данных.
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    for sheet in writer.book.worksheets:
        sheet.freeze_panes = "B2"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="284B63")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for col in sheet.columns:
            letter = get_column_letter(col[0].column)
            width = max(len(str(cell.value or "")) for cell in col)
            sheet.column_dimensions[letter].width = min(max(width + 2, 16), 70)
            for cell in col[1:]:
                if cell.data_type == "f":
                    cell.data_type = "s"   # Названия файлов и групп сохраняются как текст.
                cell.font = Font(name="Arial", size=10)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0" if isinstance(cell.value, int) else "0.00000000"
                    cell.alignment = Alignment(horizontal="right")


def safe_group_name(group):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", group.strip()).strip(" .")
    if not name or len(name) > 80:
        raise ValueError("введите название группы длиной от 1 до 80 символов")
    return name


def load_spectra(paths, say=print):
    # Прежнее чтение и проверка копий; ошибочные файлы не пропускаются.
    paths = [Path(path).resolve() for path in paths]
    if not paths or len(set(paths)) != len(paths):
        raise ValueError("нет выбранных файлов или один файл выбран повторно")
    say(f"Найдено файлов: {len(paths)}")
    spectra = []
    for path in paths:
        say(f"Обрабатывается {path.name}")
        item = read_spectrum(path, say)
        for previous in spectra:
            if (np.array_equal(item["x"], previous["x"])
                    and np.array_equal(item["y"], previous["y"])):
                raise ValueError(f"Файлы '{path.name}' и '{previous['path'].name}' "
                                 "содержат одинаковые спектры. Проверьте возможную копию")
        spectra.append(item)
    return spectra


def process_group(paths, group, signal_type, output_root, confirm=lambda text: False,
                  prepared=None, return_data=False):
    messages = []

    def say(text):
        print(text, flush=True)
        messages.append(text)

    signal_type = signal_type.upper()
    if signal_type not in Y_LABELS:
        raise ValueError("тип сигнала должен быть I, A или T")
    safe_name = safe_group_name(group)
    output_root = Path(output_root)
    reference_path = output_root / "reference_grid.json"
    if prepared is None:
        spectra = load_spectra(paths, say)
        grid, is_new = choose_grid(spectra, reference_path, signal_type, confirm, say)
    else:
        # Для двух групп сетка уже проверена по ВСЕМ исходным файлам.
        spectra, grid = prepared
        is_new = False
        say(f"Группа {group}: {len(spectra)} проверенных файлов; общая сетка двух групп.")
    normalized, file_info = [], []
    for index, item in enumerate(spectra, start=1):
        exact = np.array_equal(item["x"], grid)
        if exact:
            y = item["y"].copy()
            say(f"  {item['path'].name}: сетка совпадает, интерполяция не нужна.")
        else:
            # Только линейная интерполяция внутри измеренного диапазона.
            y = np.interp(grid, item["x"], item["y"], left=np.nan, right=np.nan)
            say(f"  {item['path'].name}: линейная интерполяция на общую сетку.")
        if not np.isfinite(y).all():
            raise ValueError(f"Файл '{item['path'].name}': ошибка интерполяции")
        norm = float(np.linalg.norm(y))
        if not np.isfinite(norm) or norm <= 0:
            raise ValueError(f"Файл '{item['path'].name}': L2-норма равна нулю или некорректна")
        # Сначала нормализуется КАЖДАЯ мышь. Среднее здесь не используется.
        normalized.append(y / norm)
        file_info.append({"Spectrum_ID": f"Mouse_{index:03d}", "File": item["path"].name,
                          "Sheet": item["sheet"], "Original_points": len(item["x"]),
                          "Original_min_cm-1": float(item["x"][0]),
                          "Original_max_cm-1": float(item["x"][-1]),
                          "Median_step_cm-1": item["step"], "L2_factor": norm,
                          "Interpolated": not exact, "Rows_sorted": item["sorted"]})
    matrix = np.vstack(normalized)   # Строка матрицы = одна мышь.
    n = matrix.shape[0]
    mean = matrix.mean(axis=0)       # Усреднение только после индивидуальной нормализации.
    sd = matrix.std(axis=0, ddof=1) if n > 1 else np.full(len(grid), np.nan)
    sem = sd / np.sqrt(n)
    say("Нормализация завершена")
    say(f"Количество спектров: {n}")
    if n == 1:
        say("При N=1 SD и SEM не определены: в Excel эти ячейки будут пустыми.")
    summary = pd.DataFrame({"Wavenumber_cm-1": grid[::-1],
                            "Mean_Normalized_Intensity": mean[::-1],
                            "SD": sd[::-1], "SEM": sem[::-1], "N": n})
    individual = pd.DataFrame(matrix[:, ::-1].T,
                              columns=[row["Spectrum_ID"] for row in file_info])
    individual.insert(0, "Wavenumber_cm-1", grid[::-1])
    metadata = pd.DataFrame({"Parameter": ["Group", "N", "Normalization", "Formula",
        "Signal_type", "Grid_min_cm-1", "Grid_max_cm-1", "Grid_points", "SD_definition",
        "SEM_definition", "Additional_preprocessing", "Python", "NumPy", "pandas",
        "matplotlib", "openpyxl"], "Value": [group, n, METHOD, "y / sqrt(sum(y**2))",
        signal_type, float(grid[0]), float(grid[-1]), len(grid), "Sample SD, ddof=1",
        "SD / sqrt(N)", "None", sys.version.split()[0], np.__version__, pd.__version__,
        matplotlib.__version__, openpyxl.__version__]})
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    run_dir = output_root / f"{safe_name}_{stamp}"
    run_dir.mkdir(parents=True, exist_ok=False)
    excel_path = run_dir / f"{safe_name}_mean_normalized.xlsx"
    png_path = run_dir / f"{safe_name}_mean_normalized.png"
    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Mean_spectrum", index=False)
        individual.to_excel(writer, sheet_name="Individual_spectra", index=False)
        metadata.to_excel(writer, sheet_name="Processing", index=False)
        pd.DataFrame(file_info).to_excel(writer, sheet_name="Processing", index=False,
                                        startrow=len(metadata) + 3)
        format_excel(writer)
    fig, ax = plt.subplots(figsize=(8, 4.5), constrained_layout=True)
    if n > 1:
        ax.fill_between(grid, mean - sd, mean + sd, color="#23638C", alpha=0.20,
                        linewidth=0, label="Mean ± SD")
    ax.plot(grid, mean, color="#174E75", linewidth=1.4, label="Mean")
    ax.set(xlabel="Wavenumber (cm⁻¹)", ylabel=Y_LABELS[signal_type],
           title=f"{group} (n={n})", xlim=(grid[-1], grid[0]))
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(labelsize=10, direction="out")
    ax.xaxis.label.set_size(12)
    ax.yaxis.label.set_size(12)
    ax.legend(frameon=False, fontsize=10)
    fig.savefig(png_path, dpi=300, facecolor="white")
    if is_new:
        reference = {"method": METHOD, "signal_type": signal_type,
                     "wavenumber_cm-1": grid.tolist()}
        reference_path.write_text(json.dumps(reference, ensure_ascii=False, indent=2),
                                  encoding="utf-8")
    say(f"Результат сохранён: {excel_path}")
    say(f"График сохранён: {png_path}")
    (run_dir / f"{safe_name}_processing_log.txt").write_text(
        "\n".join(messages), encoding="utf-8-sig")
    if return_data:
        return {"group": group, "grid": grid, "matrix": matrix, "mean": mean,
                "sd": sd, "sem": sem, "n": n, "individual": individual,
                "files": file_info, "excel": excel_path, "png": png_path,
                "signal_type": signal_type}
    return excel_path, png_path


def permutation_pvalues(cancer, control, observed_u, say):
    # Перестановки меток МЫШЕЙ, не спектральных точек. Ранги считаются один раз.
    nc, nt = len(cancer), len(control)
    total, small = nc + nt, min(nc, nt)
    ranks = stats.rankdata(np.vstack([cancer, control]), axis=0, method="average")
    possible = math.comb(total, small)
    exact = possible <= MAX_EXACT_PERMUTATIONS
    count = possible if exact else RANDOM_PERMUTATIONS
    method = "permutation_exact" if exact else f"permutation_MC_{count}"
    say(f"  Совпадающие ранги: {cancer.shape[1]} областей; {method}, перестановок {count}.")
    if exact:
        assignments = combinations(range(total), small)
    else:
        rng = np.random.default_rng(RANDOM_SEED)
        assignments = (rng.choice(total, size=small, replace=False) for _ in range(count))
        say(f"  Случайные перестановки: фиксированный seed={RANDOM_SEED}; p приближённые.")
    lower = np.zeros(cancer.shape[1], dtype=np.int64)
    upper = lower.copy()
    done = 0
    while True:
        batch = list(islice(assignments, 128))
        if not batch:
            break
        indices = np.asarray(batch, dtype=int)
        u = ranks[indices].sum(axis=1) - small * (small + 1) / 2
        if nc > nt:
            u = nc * nt - u     # Выбрана меньшая Control: переводим в U_Cancer.
        lower += (u <= observed_u).sum(axis=0)
        upper += (u >= observed_u).sum(axis=0)
        done += len(batch)
    # Двустороннее p = 2 * меньший хвост. Для Monte Carlo используется поправка +1.
    correction = 0 if exact else 1
    p = np.minimum(1.0, 2 * np.minimum(lower + correction, upper + correction)
                   / (done + correction))
    return p, method


def mann_whitney_statistics(cancer, control, say=print):
    # Матрицы: строки = мыши; столбцы = значения для проверяемых областей.
    cancer, control = np.asarray(cancer, float), np.asarray(control, float)
    if (cancer.ndim != 2 or control.ndim != 2 or cancer.shape[1] != control.shape[1]
            or cancer.shape[1] == 0 or min(len(cancer), len(control)) < 2):
        raise ValueError("для сравнения нужны две матрицы, минимум две мыши в каждой группе")
    if not np.isfinite(cancer).all() or not np.isfinite(control).all():
        raise ValueError("в индивидуальных спектрах обнаружены NaN или бесконечность")
    nc, nt = len(cancer), len(control)
    p, u = np.empty(cancer.shape[1]), np.empty(cancer.shape[1])
    methods = np.empty(cancer.shape[1], dtype=object)
    combined = np.sort(np.vstack([cancer, control]), axis=0)
    ties = np.any(np.diff(combined, axis=0) == 0, axis=0)
    small = min(nc, nt) <= 8
    exact_columns = ~ties if small else np.zeros(len(p), dtype=bool)
    for mask, method in ((exact_columns, "exact"), (~exact_columns, "asymptotic")):
        if mask.any():
            result = stats.mannwhitneyu(cancer[:, mask], control[:, mask], axis=0,
                alternative="two-sided", method=method, use_continuity=True)
            u[mask], p[mask], methods[mask] = result.statistic, result.pvalue, method
    if small and ties.any():
        p[ties], methods[ties] = permutation_pvalues(
            cancer[:, ties], control[:, ties], u[ties], say)
    if not np.isfinite(p).all() or np.any((p < 0) | (p > 1)):
        raise ValueError("не удалось корректно рассчитать p-values")
    # Одна BH-поправка по ВСЕМ проверенным областям, без отбора по исходному p.
    q = false_discovery_control(p, method="bh")
    # SciPy возвращает U первой группы; первой передана Cancer.
    # U_Cancer = число пар Cancer > Control + половина числа равных пар.
    # Эффект = (число выигранных пар - число проигранных пар) / (nc * nt).
    effect = 2 * u / (nc * nt) - 1
    say(f"Mann–Whitney: {len(p)} областей; BH-FDR: {int((q < FDR_LEVEL).sum())} "
        f"областей с q < {FDR_LEVEL}.")
    return {"u": u, "p": p, "q": q, "effect": effect, "methods": methods}


def confidence_interval(group):
    # Точечный 95% интервал среднего по Стьюденту; не одновременная полоса спектра.
    half = stats.t.ppf(0.975, df=group["n"] - 1) * group["sem"]
    return group["mean"] - half, group["mean"] + half


def mean_signal_in_region(grid, matrix, lower, upper):
    """Средний сигнал каждой мыши: интеграл по области / её точная ширина."""
    grid, matrix = np.asarray(grid, float), np.asarray(matrix, float)
    if (grid.ndim != 1 or len(grid) < 2 or not np.isfinite(grid).all()
            or np.any(np.diff(grid) <= 0) or matrix.ndim != 2
            or matrix.shape[1] != len(grid) or not np.isfinite(matrix).all()):
        raise ValueError("нужны возрастающая сетка и конечные индивидуальные спектры")
    if not np.isfinite([lower, upper]).all() or lower >= upper:
        raise ValueError("некорректные границы спектральной области")
    if grid[0] > lower or grid[-1] < upper:
        raise ValueError("общая сетка не покрывает область полностью; экстраполяция запрещена")
    # Даже когда точных границ нет в сетке, интеграл имеет заданные пределы.
    inside = (grid > lower) & (grid < upper)
    region_grid = np.r_[lower, grid[inside], upper]
    left = np.array([np.interp(lower, grid, row) for row in matrix])
    right = np.array([np.interp(upper, grid, row) for row in matrix])
    region_matrix = np.column_stack([left, matrix[:, inside], right])
    return trapezoid(region_matrix, x=region_grid, axis=1) / (upper - lower)


def regional_statistics(grid, cancer, control, say=print):
    """Одна независимая выборка мышей на группу, один тест на заданную область."""
    grid = np.asarray(grid, float)
    cancer, control = np.asarray(cancer, float), np.asarray(control, float)
    if (grid.ndim != 1 or len(grid) < 2 or not np.isfinite(grid).all()
            or np.any(np.diff(grid) <= 0)):
        raise ValueError("нужна конечная возрастающая сетка волновых чисел")
    if (cancer.ndim != 2 or control.ndim != 2
            or cancer.shape[1] != len(grid) or control.shape[1] != len(grid)
            or min(len(cancer), len(control)) < 2):
        raise ValueError("нужны индивидуальные спектры на общей сетке и минимум две мыши в группе")
    if not np.isfinite(cancer).all() or not np.isfinite(control).all():
        raise ValueError("в индивидуальных спектрах обнаружены NaN или бесконечность")
    tested = np.array([grid[0] <= lower and grid[-1] >= upper
                       for _, lower, upper in SPECTRAL_REGIONS], dtype=bool)
    if not tested.any():
        raise ValueError("общая сетка не покрывает полностью ни одну из девяти заданных областей")
    cancer_values = np.full((len(cancer), len(SPECTRAL_REGIONS)), np.nan)
    control_values = np.full((len(control), len(SPECTRAL_REGIONS)), np.nan)
    for index, (name, lower, upper) in enumerate(SPECTRAL_REGIONS):
        if not tested[index]:
            say(f"Не проверена: {name}, {upper:g}–{lower:g} см⁻¹ — нет полного покрытия.")
            continue
        cancer_values[:, index] = mean_signal_in_region(grid, cancer, lower, upper)
        control_values[:, index] = mean_signal_in_region(grid, control, lower, upper)
    say("Показатель области: интеграл индивидуального L2-нормированного спектра / ширина области.")
    calculated = mann_whitney_statistics(cancer_values[:, tested], control_values[:, tested], say)
    full = {}
    for key in ("u", "p", "q", "effect"):
        full[key] = np.full(len(SPECTRAL_REGIONS), np.nan)
        full[key][tested] = calculated[key]
    methods = np.full(len(SPECTRAL_REGIONS), "not_tested", dtype=object)
    methods[tested] = calculated["methods"]
    rows = []
    for index, (name, lower, upper) in enumerate(SPECTRAL_REGIONS):
        row = {"Region_ID": f"R{index + 1}", "Region": name,
            "Lower_cm-1": lower, "Upper_cm-1": upper, "Width_cm-1": upper - lower,
            "Tested": bool(tested[index]),
            "Coverage_status": "Полное покрытие" if tested[index] else "Нет полного покрытия",
            "Cancer_N": len(cancer), "Control_N": len(control)}
        for label, values in (("Cancer", cancer_values), ("Control", control_values)):
            sample = values[:, index]
            row[f"{label}_mean_region_signal"] = float(sample.mean()) if tested[index] else np.nan
            row[f"{label}_median_region_signal"] = float(np.median(sample)) if tested[index] else np.nan
            row[f"{label}_SD_region_signal"] = float(sample.std(ddof=1)) if tested[index] else np.nan
        row.update({"Difference_Cancer_minus_Control":
                row["Cancer_mean_region_signal"] - row["Control_mean_region_signal"],
            "Mann_Whitney_U": full["u"][index], "p_value": full["p"][index],
            "q_value_BH_FDR": full["q"][index], "Effect_size": full["effect"][index],
            "Significant_q_0.05": bool(full["q"][index] < FDR_LEVEL) if tested[index] else None,
            "P_value_method": methods[index],
            "Direction_by_rank": ("Cancer > Control" if full["effect"][index] > 0 else
                                  "Cancer < Control" if full["effect"][index] < 0 else
                                  "Нет рангового направления" if tested[index] else "Не проверена")})
        row["Conclusion"] = ("Не проверена: спектры не покрывают область целиком" if not tested[index] else
            "Обнаружены статистически значимые различия" if row["Significant_q_0.05"] else
            "Статистически значимые различия не обнаружены")
        rows.append(row)
        if tested[index]:
            say(f"{name} ({upper:g}–{lower:g} см⁻¹): p={full['p'][index]:.6g}, "
                f"q={full['q'][index]:.6g}. {row['Conclusion']}.")
    return {"statistics": pd.DataFrame(rows), "cancer_values": cancer_values,
            "control_values": control_values}


def make_summary_table(statistics):
    """Компактная таблица для отчёта; точные p и q остаются на подробном листе."""
    def report_p(value, tested):
        if not tested:
            return "Не проверена"
        return "< 0.0001" if value < 0.0001 else float(value)

    rows = []
    for row in statistics.to_dict("records"):
        rows.append({"Region": row["Region"],
            "Range, cm⁻¹": f"{row['Upper_cm-1']:g}–{row['Lower_cm-1']:g}",
            "Median control": row["Control_median_region_signal"],
            "Median cancer": row["Cancer_median_region_signal"],
            "p-MWU": report_p(row["p_value"], row["Tested"]),
            "p-BH": report_p(row["q_value_BH_FDR"], row["Tested"]),
            "r": row["Effect_size"]})
    return pd.DataFrame(rows)


def format_summary_sheet(sheet):
    # Белый фон, горизонтальные линии и центрированные числа, как в образце.
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.page import PageMargins

    sheet.freeze_panes = "C2"
    widths = {"A": 54, "B": 20, "C": 19, "D": 19, "E": 15, "F": 15, "G": 12}
    line = Side(style="thin", color="505050")
    for letter, width in widths.items():
        sheet.column_dimensions[letter].width = width
    sheet.row_dimensions[1].height = 34
    for header in sheet[1]:
        header.font = Font(name="Times New Roman", size=12, bold=True,
                           italic=header.column in (5, 6, 7), color="222222")
        header.fill = PatternFill(fill_type=None)
        header.alignment = Alignment(horizontal="left" if header.column == 1 else "center",
                                     vertical="center", wrap_text=True)
        header.border = Border(top=line, bottom=line)
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 40
        for cell in row:
            cell.font = Font(name="Times New Roman", size=12, color="222222")
            cell.alignment = Alignment(horizontal="left" if cell.column == 1 else "center",
                                       vertical="center", wrap_text=True)
            if cell.column in (3, 4, 5, 6):
                cell.number_format = "0.0000"
            elif cell.column == 7:
                cell.number_format = "0.000"
            if cell.row == sheet.max_row:
                cell.border = Border(bottom=line)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_options.horizontalCentered = True
    sheet.print_title_rows = "1:1"
    sheet.print_area = f"A1:G{sheet.max_row}"
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4)


def save_comparison_figures(control, cancer, difference, regions, band, folder, stem):
    from matplotlib.patches import Patch

    grid = control["grid"]
    blue, red = "#23638C", "#B5453A"
    fig, (ax, strip) = plt.subplots(2, 1, figsize=(10, 8), sharex=True,
        gridspec_kw={"height_ratios": [5, 4]}, constrained_layout=True)
    for group, color in ((control, blue), (cancer, red)):
        low, high = (group["mean"] - group["sd"], group["mean"] + group["sd"])
        if band == "CI":
            low, high = confidence_interval(group)
        ax.fill_between(grid, low, high, color=color, alpha=0.18, linewidth=0)
        ax.plot(grid, group["mean"], color=color, linewidth=1.4,
                label=f"{group['group']} (n={group['n']})")
    ax.set(ylabel=Y_LABELS[control["signal_type"]],
           title="Mean ± SD" if band == "SD" else "Mean ± pointwise 95% CI")
    ax.legend(frameon=False, fontsize=10)
    # Отдельная строка для каждой гипотезы: перекрывающиеся области не сливаются.
    labels = []
    for index, row in enumerate(regions.to_dict("records")):
        labels.append(f"{row['Region_ID']}: {row['Upper_cm-1']:g}–{row['Lower_cm-1']:g}")
        left, right = max(grid[0], row["Lower_cm-1"]), min(grid[-1], row["Upper_cm-1"])
        color = (red if row["Effect_size"] > 0 else blue) if row["Significant_q_0.05"] else "0.82"
        if right > left:
            strip.broken_barh([(left, right - left)], (index - 0.32, 0.64),
                facecolors=color, edgecolors="0.5" if not row["Tested"] else "none",
                hatch="///" if not row["Tested"] else None, alpha=0.8)
        if not row["Tested"]:
            strip.text((grid[0] + grid[-1]) / 2, index, "not tested: incomplete coverage",
                       ha="center", va="center", fontsize=8)
    strip.set(xlabel="Wavenumber (cm⁻¹)", xlim=(grid[-1], grid[0]),
              ylim=(len(regions) - 0.5, -0.5), title="Tests of mean signal in predefined regions")
    strip.set_yticks(np.arange(len(regions)), labels=labels)
    strip.legend(handles=[Patch(color=red, label=f"{cancer['group']} higher: q < 0.05"),
        Patch(color=blue, label=f"{control['group']} higher: q < 0.05"),
        Patch(color="0.82", label="No significant difference"),
        Patch(facecolor="0.82", edgecolor="0.5", hatch="///", label="Not tested")],
        loc="upper center", bbox_to_anchor=(0.5, -0.24), ncol=2, frameon=False, fontsize=8)
    for axis in (ax, strip):
        axis.spines[["top", "right"]].set_visible(False)
        axis.tick_params(labelsize=10, direction="out")
        axis.xaxis.label.set_size(12)
        axis.yaxis.label.set_size(12 if axis is ax else 9)
    strip.spines["left"].set_visible(False)
    paths = []
    for extension in ("png", "svg"):
        path = folder / f"{stem}_comparison.{extension}"
        fig.savefig(path, dpi=300, facecolor="white")
        paths.append(path)
    fig_diff, ax_diff = plt.subplots(figsize=(8, 4.5), constrained_layout=True)
    ax_diff.axhline(0, color="0.4", linewidth=0.8, linestyle="--")
    ax_diff.plot(grid, difference, color="#333333", linewidth=1.3)
    ax_diff.set(xlabel="Wavenumber (cm⁻¹)", ylabel="Δ " + Y_LABELS[control["signal_type"]],
        title=f"{cancer['group']} − {control['group']}", xlim=(grid[-1], grid[0]))
    ax_diff.spines[["top", "right"]].set_visible(False)
    ax_diff.tick_params(labelsize=10, direction="out")
    ax_diff.xaxis.label.set_size(12)
    ax_diff.yaxis.label.set_size(12)
    for extension in ("png", "svg"):
        path = folder / f"{stem}_difference.{extension}"
        fig_diff.savefig(path, dpi=300, facecolor="white")
        paths.append(path)
    return paths


def compare_groups(control_paths, cancer_paths, control_name, cancer_name, signal_type,
                   output_root, band="SD", confirm=lambda text: False):
    messages = []

    def say(text):
        print(text, flush=True)
        messages.append(text)

    control_safe, cancer_safe = safe_group_name(control_name), safe_group_name(cancer_name)
    if control_safe.casefold() == cancer_safe.casefold():
        raise ValueError("названия двух групп должны различаться")
    signal_type, band = signal_type.strip().upper(), band.strip().upper()
    if signal_type not in Y_LABELS or band not in ("SD", "CI"):
        raise ValueError("тип сигнала: I/A/T; область на общем графике: SD или CI")
    say(f"Загрузка первой группы: {control_name}")
    control_raw = load_spectra(control_paths, say)
    say(f"Загрузка второй группы: {cancer_name}")
    cancer_raw = load_spectra(cancer_paths, say)
    if min(len(control_raw), len(cancer_raw)) < 2:
        raise ValueError("для сравнения и 95% CI нужны минимум две мыши в каждой группе")
    for cancer_item in cancer_raw:
        for control_item in control_raw:
            if (cancer_item["path"] == control_item["path"] or
                    (np.array_equal(cancer_item["x"], control_item["x"]) and
                     np.array_equal(cancer_item["y"], control_item["y"]))):
                raise ValueError(f"Файлы '{cancer_item['path']}' и '{control_item['path']}' "
                    "повторяются между группами. Для независимого теста нужны разные мыши")
    output_root = Path(output_root)
    reference_path = output_root / "reference_grid.json"
    # Прежний выбор сетки, но одновременно для всех мышей обеих групп.
    grid, is_new = choose_grid(control_raw + cancer_raw, reference_path,
                              signal_type, confirm, say)
    if not any(grid[0] <= lower and grid[-1] >= upper for _, lower, upper in SPECTRAL_REGIONS):
        raise ValueError("общая сетка не покрывает полностью ни одну из девяти заданных областей")
    folder = output_root / ("Comparison_" + datetime.now().strftime("%Y%m%d_%H%M%S_%f"))
    folder.mkdir(parents=True, exist_ok=False)
    control = process_group(control_paths, control_name, signal_type, folder,
                            prepared=(control_raw, grid), return_data=True)
    cancer = process_group(cancer_paths, cancer_name, signal_type, folder,
                           prepared=(cancer_raw, grid), return_data=True)
    say("Статистическая единица — отдельная мышь. Пропущено файлов: 0.")
    calculated = regional_statistics(grid, cancer["matrix"], control["matrix"], say)
    difference = cancer["mean"] - control["mean"]
    statistics = calculated["statistics"]
    summary_table = make_summary_table(statistics)
    significant = statistics.loc[statistics["Significant_q_0.05"].eq(True)].copy()
    value_rows = []
    for index, (name, lower, upper) in enumerate(SPECTRAL_REGIONS):
        for role, group, values in (("Cancer", cancer, calculated["cancer_values"]),
                                   ("Control", control, calculated["control_values"])):
            for mouse_index, info in enumerate(group["files"]):
                value_rows.append({"Region_ID": f"R{index + 1}", "Region": name,
                    "Lower_cm-1": lower, "Upper_cm-1": upper, "Role": role,
                    "Group": group["group"], "Spectrum_ID": info["Spectrum_ID"],
                    "File": info["File"], "Mean_normalized_region_signal": values[mouse_index, index],
                    "Coverage_status": statistics.iloc[index]["Coverage_status"]})
    region_values = pd.DataFrame(value_rows)
    means = {"Wavenumber_cm-1": grid}
    for label, group in (("Cancer", cancer), ("Control", control)):
        low, high = confidence_interval(group)
        means.update({f"{label}_mean": group["mean"], f"{label}_SD": group["sd"],
            f"{label}_SEM": group["sem"], f"{label}_CI95_lower": low,
            f"{label}_CI95_upper": high, f"{label}_N": group["n"]})
    means["Difference_Cancer_minus_Control"] = difference
    means = pd.DataFrame(means).iloc[::-1].reset_index(drop=True)
    stem = f"{control_safe[:40]}_vs_{cancer_safe[:40]}"
    excel_path = folder / f"{stem}_statistics.xlsx"
    settings = {"Control_label": control_name, "Cancer_label": cancer_name,
        "Control_N": control["n"], "Cancer_N": cancer["n"], "Skipped_files": 0,
        "Normalization": METHOD, "Formula": "y / sqrt(sum(y**2)), each mouse separately",
        "Signal_type": signal_type, "Additional_preprocessing": "None",
        "Grid_min_cm-1": grid[0], "Grid_max_cm-1": grid[-1], "Grid_points": len(grid),
        "Test": "Mann-Whitney U on one regional mean signal per mouse; two-sided; Cancer first",
        "Region_metric": "Trapezoidal integral of individual L2-normalized signal / (upper-lower)",
        "Region_boundaries": "Exact requested limits; linear interpolation at boundaries; no extrapolation",
        "Region_coverage": "Full coverage required; incomplete regions not tested and excluded from FDR",
        "Requested_regions": len(SPECTRAL_REGIONS), "Tested_regions": int(statistics["Tested"].sum()),
        "P_method": "Small n: exact without ties, permutation with ties; otherwise asymptotic",
        "FDR": "One Benjamini-Hochberg correction over all fully covered predefined regions; q < 0.05",
        "Effect": "Rank-biserial = 2 * U_Cancer / (N_Cancer * N_Control) - 1",
        "Effect_direction": "Positive: Cancer ranks higher; negative: Control ranks higher",
        "Summary_medians": "Median of individual regional mean signals in each group, not median across wavenumbers",
        "Summary_r": "Rank-biserial effect (same as detailed Effect_size); not Z / sqrt(N)",
        "Summary_p_display": "p < 0.0001 displayed as < 0.0001; unrounded numeric p/q in Region_statistics",
        "CI95": "Pointwise mean +/- t(0.975, N-1) * SD / sqrt(N); normality assumption",
        "Regions": "Nine predefined regions in requested order; overlapping regions kept as separate tests",
        "Interpretation": "Compares regional mean normalized signal; not an omnibus test of spectral shape",
        "Not_significant": "No statistically significant difference detected; does not establish equivalence",
        "Region_colors": "Red: Cancer ranks higher; blue: Control ranks higher; gray: not significant",
        "Descriptive_spectra": "Full-grid means, SD and pointwise CI are descriptive; no pointwise hypothesis tests",
        "Comparison_band": band, "Permutation_exact_limit": MAX_EXACT_PERMUTATIONS,
        "Permutation_MC_count": RANDOM_PERMUTATIONS, "Permutation_seed": RANDOM_SEED,
        "Python": sys.version.split()[0], "SciPy": scipy.__version__}
    metadata = pd.DataFrame(settings.items(), columns=["Parameter", "Value"])
    mapping = pd.DataFrame([dict(Role=role, Group=group["group"], **row)
        for role, group in (("Cancer", cancer), ("Control", control)) for row in group["files"]])
    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        for table, sheet in ((summary_table, "Summary_table"),
                (statistics, "Region_statistics"), (significant, "Significant_regions"),
                (region_values, "Region_values"),
                (cancer["individual"], "Cancer_individual_spectra"),
                (control["individual"], "Control_individual_spectra"), (means, "Group_means"),
                (metadata, "Processing")):
            table.to_excel(writer, sheet_name=sheet, index=False)
        mapping.to_excel(writer, sheet_name="Processing", index=False, startrow=len(metadata) + 3)
        format_excel(writer)
        format_summary_sheet(writer.book["Summary_table"])
        # Малые p и q видны в Excel, без отображения их как округлённых нулей.
        for sheet_name in ("Region_statistics", "Significant_regions"):
            sheet = writer.book[sheet_name]
            for header in sheet[1]:
                if header.value in ("p_value", "q_value_BH_FDR"):
                    for cells in sheet.iter_cols(min_col=header.column, max_col=header.column, min_row=2):
                        for cell in cells:
                            cell.number_format = "0.0000E+00"
    figures = save_comparison_figures(control, cancer, difference, statistics, band, folder, stem)
    if is_new:
        reference = {"method": METHOD, "signal_type": signal_type, "wavenumber_cm-1": grid.tolist()}
        reference_path.write_text(json.dumps(reference, ensure_ascii=False, indent=2), encoding="utf-8")
    say(f"{control_name}: {control['n']} spectra successfully processed")
    say(f"{cancer_name}: {cancer['n']} spectra successfully processed")
    say(f"Значимых областей: {len(significant)} из {int(statistics['Tested'].sum())} проверенных.")
    say("Статистические тесты отдельных волновых чисел не проводились.")
    say("Тесты сравнивают средний нормированный сигнал области, а не все возможные различия формы спектра.")
    say("Перекрывающиеся области не являются независимыми подтверждениями различий.")
    say(f"Результат сохранён: {excel_path}")
    for path in figures:
        say(f"График сохранён: {path}")
    (folder / "comparison_log.txt").write_text("\n".join(messages), encoding="utf-8-sig")
    return {"excel": excel_path, "figures": figures, "folder": folder,
            "control": control, "cancer": cancer, "statistics": statistics,
            "summary_table": summary_table, "region_values": region_values,
            "significant_regions": significant}


def main():
    root = tk.Tk()
    root.withdraw()
    output_root = Path(__file__).resolve().parent / "IR_results"
    try:
        comparison = messagebox.askyesno("Режим обработки",
            "Сравнить две независимые группы?\n\nДа — Control и Breast cancer.\n"
            "Нет — прежняя обработка одной группы.\n\n"
            "Сравнение выполняется по среднему нормированному сигналу\n"
            "в девяти заданных спектральных областях.\n"
            "Каждый файл должен соответствовать одной отдельной мыши.", parent=root)
        files = filedialog.askopenfilenames(parent=root,
            title="Выберите Excel-файлы Control" if comparison else "Выберите Excel-файлы одной группы",
            filetypes=[("Excel", "*.xlsx *.xls"), ("Все файлы", "*.*")])
        if not files:
            print("Файлы не выбраны. Программа завершена.")
            return
        group = simpledialog.askstring("Название первой группы", "Введите название группы:",
                                       initialvalue="Control", parent=root)
        if group is None:
            return
        if comparison:
            cancer_files = filedialog.askopenfilenames(parent=root,
                title="Выберите Excel-файлы Breast cancer",
                filetypes=[("Excel", "*.xlsx *.xls"), ("Все файлы", "*.*")])
            if not cancer_files:
                print("Вторая группа не выбрана. Сравнение отменено.")
                return
            cancer_name = simpledialog.askstring("Название второй группы",
                "Введите название группы РМЖ:", initialvalue="Breast_cancer", parent=root)
            if cancer_name is None:
                return
        signal_type = simpledialog.askstring("Тип исходного сигнала",
            "Введите одну латинскую букву:\nA — поглощение (Absorbance)\n"
            "T — пропускание (Transmittance)\nI — интенсивность или неизвестный сигнал\n\n"
            "Если тип не указан в экспорте, выберите I.\n"
            "Выбор задаёт подпись оси. Пересчёта пропускания в поглощение нет.",
            initialvalue="I", parent=root)
        if signal_type is None:
            return
        confirm = lambda text: messagebox.askyesno("Общая область спектров", text, parent=root)
        if comparison:
            band = simpledialog.askstring("Область на общем графике",
                "Введите SD или CI латинскими буквами:\n"
                "SD — среднее ± стандартное отклонение (по умолчанию).\n"
                "CI — среднее ± точечный 95% интервал среднего.\n\n"
                "При малом N интервал CI зависит от формы распределения.\n"
                "Выбор области не меняет тесты и BH-FDR.", initialvalue="SD", parent=root)
            if band is None:
                return
            result = compare_groups(files, cancer_files, group, cancer_name,
                                    signal_type, output_root, band, confirm)
            destination = result["folder"]
        else:
            excel_path, png_path = process_group(files, group, signal_type.strip().upper(),
                                               output_root, confirm=confirm)
            destination = excel_path.parent
        details = ("Таблица по образцу: лист Summary_table.\n"
                   "Выводы по девяти областям: лист Region_statistics.\n"
                   "Показатель: средний нормированный сигнал внутри области.\n"
                   if comparison else "")
        messagebox.showinfo("Готово", f"Результаты сохранены в папке:\n{destination}\n\n"
                            + details + "После нажатия ОК откроются графики.", parent=root)
    except Exception as exc:
        print(f"ОШИБКА: {exc}")
        try:
            output_root.mkdir(parents=True, exist_ok=True)
            error_path = output_root / ("error_" + datetime.now().strftime("%Y%m%d_%H%M%S") + ".txt")
            error_path.write_text(traceback.format_exc(), encoding="utf-8-sig")
            print(f"Описание ошибки сохранено: {error_path}")
        except OSError:
            pass
        messagebox.showerror("Ошибка обработки", str(exc), parent=root)
        return
    finally:
        root.destroy()
    plt.show()


if __name__ == "__main__":
    main()
