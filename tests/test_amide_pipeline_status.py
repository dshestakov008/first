"""Keep diagnostic decompositions separate from accepted animal statistics."""

import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from amide_analysis import __version__
from amide_analysis.config import default_config
from amide_analysis.fitting import STRUCTURES
from amide_analysis.pipeline import run_analysis
from amide_analysis.reports import MAIN_SHEETS


class PipelineStatusTests(unittest.TestCase):
    def _run(self, directory, *, confirmed=False, numerical=True, failed_variant=False,
             invalid_source=False):
        input_dir = Path(directory) / "spectra"
        for group in ("Control", "Cancer"):
            group_dir = input_dir / group
            group_dir.mkdir(parents=True)
            x = np.arange(1580.0, 1720.0, 1.92905)
            y = 0.1 + 0.0001 * (x - 1600.0) + np.exp(-4 * np.log(2) * ((x - 1655.0) / 18.0) ** 2)
            source = pd.DataFrame({"wavenumber": x, "absorbance": y})
            if invalid_source and group == "Cancer":
                source["absorbance"] = source["absorbance"].astype(object)
                source.loc[5, "absorbance"] = "invalid signal"
            source.to_csv(group_dir / "mouse.csv", index=False)
        config = default_config()
        config["resolution_confirmed"] = confirmed
        fractions = dict(zip(STRUCTURES, [10.0, 60.0, 20.0, 5.0, 5.0]))

        def fitted(processed, config, seed=None, **kwargs):
            accepted = bool(numerical and confirmed)
            return {
                "status": "Reliable fit" if accepted else "Unreliable fit",
                "warnings": [] if numerical else ["Bootstrap areas are not stable"],
                "fit": processed["corrected"].copy(),
                "residuals": np.zeros_like(processed["x"]),
                "metrics": {"components": 1, "r_squared": .999,
                            "numerical_fit_reliable": numerical, "quantitative_accepted": accepted},
                "numerical_fit_reliable": numerical,
                "quantitative_accepted": accepted,
                "fit_rejection_codes": [] if numerical else ["BOOTSTRAP_AREA_INSTABILITY"],
                "protocol_blockers": [] if confirmed else ["INSTRUMENT_RESOLUTION_UNCONFIRMED"],
                "percentages": fractions.copy() if accepted else {key: np.nan for key in STRUCTURES},
                "diagnostic_percentages": fractions.copy(),
                "components": [{"accepted": accepted, "area_pct": 100.0 if accepted else np.nan}],
                "candidates": [], "model_comparison": [], "sensitivity": [],
                "bootstrap_diagnostics": [
                    {"iteration": 1, "status": "accepted", "optimization_converged": True},
                    {"iteration": 2, "status": "quality_rejected", "optimization_converged": True,
                     "reason": "area instability"},
                ],
            }

        messages = []
        patches = [patch("amide_analysis.pipeline.analyze_band", side_effect=fitted),
                   patch("amide_analysis.pipeline.plot_spectrum", return_value=[]),
                   patch("amide_analysis.pipeline.plot_groups", return_value=[])]
        if failed_variant:
            patches.append(patch("amide_analysis.pipeline._processing_variants",
                                 return_value=([], ["Alternative baseline could not be processed"])))
        with patches[0], patches[1], patches[2]:
            if failed_variant:
                with patches[3]:
                    result = run_analysis(input_dir, Path(directory) / "results", config=config,
                                          progress=messages.append)
            else:
                result = run_analysis(input_dir, Path(directory) / "results", config=config,
                                      progress=messages.append)
        return result, messages, fractions

    def test_unconfirmed_resolution_preserves_diagnostics_without_inference(self):
        with tempfile.TemporaryDirectory() as directory:
            result, messages, fractions = self._run(directory)
            quality = result["tables"]["Fit_Quality"]
            self.assertTrue(quality["numerical_fit_reliable"].all())
            self.assertTrue(quality["analysis_seconds"].ge(0).all())
            self.assertFalse(quality["quantitative_accepted"].any())
            self.assertEqual(quality["fit_rejection_codes"].tolist(), ["[]", "[]"])
            self.assertEqual(json.loads(quality["protocol_blockers"].iloc[0]),
                             ["INSTRUMENT_RESOLUTION_UNCONFIRMED"])
            diagnostic = result["tables"]["Diagnostic_Structure"]
            for name, value in fractions.items():
                np.testing.assert_allclose(diagnostic[name], value)
            self.assertTrue(diagnostic["percentage_basis"].str.startswith("Diagnostic only").all())
            animals = result["tables"]["Secondary_Structure"]
            self.assertTrue(animals[list(STRUCTURES)].isna().all().all())
            self.assertFalse(animals["status"].eq("Reliable fit").any())
            tests = result["tables"]["Statistical_Tests"]
            self.assertFalse(np.isfinite(tests["p_value"]).any())
            self.assertEqual(len(result["tables"]["Bootstrap_Diagnostics"]), 4)
            book = load_workbook(result["excel"], read_only=True)
            self.assertEqual(book.sheetnames[:10], list(MAIN_SHEETS))
            self.assertIn("Diagnostic_Structure", book.sheetnames)
            self.assertIn("Bootstrap_Diagnostics", book.sheetnames)
            book.close()
            for name in ("Diagnostic_Structure", "Bootstrap_Diagnostics", "Fit_Quality"):
                self.assertTrue((result["output_dir"] / f"{name}.csv").is_file())
            settings = result["tables"]["Processing_Settings"].set_index("key")
            self.assertEqual(settings.at["software_version", "value"], __version__)
            log = "\n".join(messages)
            self.assertIn("численно устойчивых разложений: 2/2", log)
            self.assertIn("количественному сравнению: 0/2", log)
            self.assertIn(str(result["excel"].resolve()), log)
            self.assertIn(str(result["output_dir"].resolve()), messages[0])
            self.assertIn("ограничение протокола", log)

    def test_genuine_instability_is_distinct_from_resolution_blocker(self):
        with tempfile.TemporaryDirectory() as directory:
            result, messages, fractions = self._run(directory, numerical=False)
            quality = result["tables"]["Fit_Quality"]
            self.assertFalse(quality["numerical_fit_reliable"].any())
            self.assertEqual(json.loads(quality["fit_rejection_codes"].iloc[0]),
                             ["BOOTSTRAP_AREA_INSTABILITY"])
            self.assertEqual(json.loads(quality["protocol_blockers"].iloc[0]),
                             ["INSTRUMENT_RESOLUTION_UNCONFIRMED"])
            self.assertEqual(result["tables"]["Diagnostic_Structure"]["alpha_helix"].iloc[0],
                             fractions["alpha_helix"])
            self.assertIn("численно устойчивых разложений: 0/2", "\n".join(messages))

    def test_confirmed_resolution_accepts_only_numerically_stable_fits(self):
        with tempfile.TemporaryDirectory() as directory:
            result, messages, fractions = self._run(directory, confirmed=True)
            self.assertTrue(result["tables"]["Fit_Quality"]["quantitative_accepted"].all())
            self.assertTrue(result["tables"]["Secondary_Structure"]["status"].eq("Reliable fit").all())
            np.testing.assert_allclose(result["tables"]["Secondary_Structure"]["alpha_helix"],
                                       fractions["alpha_helix"])
            # One animal in each group remains insufficient for an inferential test.
            self.assertFalse(np.isfinite(result["tables"]["Statistical_Tests"]["p_value"]).any())
            self.assertIn("количественному сравнению: 2/2", "\n".join(messages))

    def test_unavailable_sensitivity_variant_blocks_acceptance_but_keeps_diagnostic_areas(self):
        with tempfile.TemporaryDirectory() as directory:
            result, messages, fractions = self._run(directory, confirmed=True, failed_variant=True)
            quality = result["tables"]["Fit_Quality"]
            self.assertFalse(quality["numerical_fit_reliable"].any())
            self.assertFalse(quality["quantitative_accepted"].any())
            self.assertIn("SENSITIVITY_VARIANT_UNAVAILABLE",
                          json.loads(quality["fit_rejection_codes"].iloc[0]))
            self.assertEqual(result["tables"]["Diagnostic_Structure"]["alpha_helix"].iloc[0],
                             fractions["alpha_helix"])
            self.assertTrue(result["tables"]["Secondary_Structure"][list(STRUCTURES)].isna().all().all())

    def test_bad_source_preserves_traceback_and_exports_other_spectrum(self):
        with tempfile.TemporaryDirectory() as directory:
            result, messages, fractions = self._run(directory, invalid_source=True)
            quality = result["tables"]["Fit_Quality"].set_index("group")
            self.assertEqual(quality.at["Cancer", "status"], "Failed")
            self.assertIn("SPECTRUM_PROCESSING_FAILED",
                          json.loads(quality.at["Cancer", "fit_rejection_codes"]))
            self.assertTrue(quality.at["Control", "numerical_fit_reliable"])
            diagnostic = result["tables"]["Diagnostic_Structure"].set_index("group")
            self.assertTrue(diagnostic.loc["Cancer", list(STRUCTURES)].isna().all())
            self.assertEqual(diagnostic.at["Control", "alpha_helix"], fractions["alpha_helix"])
            errors = list(result["output_dir"].glob("error_*.txt"))
            self.assertEqual(len(errors), 1)
            self.assertIn("Traceback", errors[0].read_text(encoding="utf-8"))
            self.assertTrue(result["excel"].is_file())
            self.assertIn("Прочитано спектров: 1/2", "\n".join(messages))


if __name__ == "__main__":
    unittest.main()
