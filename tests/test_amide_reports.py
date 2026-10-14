"""Meaningful export and figure checks without a GUI or real animal data."""

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from amide_analysis import plotting, reports
import matplotlib.pyplot as plt


def _config():
    return {
        "seed": 10,
        "plots": {"control_group": "Control", "figure_size": [6, 4]},
        "statistics": {"alpha": .05},
        "spectral_resolution_cm1": 4,
        "resolution_confirmed": False,
    }


class ExportTests(unittest.TestCase):
    def test_exact_main_sheets_csv_numeric_precision_and_safe_text(self):
        value = .000000123456789012345
        metadata = pd.DataFrame({"animal_id": pd.Series(["=2+3", "normal"], dtype="category"),
                                 "file": ["\t@command", "path.csv"]})
        tests = pd.DataFrame({"q_value": [value], "effect": [-.725],
                              "status": ["ok"]})
        with tempfile.TemporaryDirectory() as directory:
            result = reports.export_results(Path(directory), {
                "Metadata": metadata, "Statistical_Tests": tests,
                "Unreliable_Spectra": pd.DataFrame({"reason": ["unconfirmed resolution"]}),
            }, _config())
            book = load_workbook(result, data_only=False)
            self.assertEqual(book.sheetnames[:10], list(reports.MAIN_SHEETS))
            self.assertEqual(book["Metadata"]["A2"].value, "'=2+3")
            self.assertEqual(book["Metadata"]["A2"].data_type, "s")
            self.assertEqual(book["Metadata"]["B2"].value, "'\t@command")
            self.assertAlmostEqual(book["Statistical_Tests"]["A2"].value, value, delta=1e-21)
            self.assertEqual(book["Statistical_Tests"]["B2"].value, -.725)
            for name in reports.MAIN_SHEETS:
                self.assertTrue((Path(directory) / f"{name}.csv").exists())
            csv = pd.read_csv(Path(directory) / "Statistical_Tests.csv")
            self.assertAlmostEqual(csv["q_value"].iloc[0], value, delta=1e-21)
            settings = pd.read_csv(Path(directory) / "Processing_Settings.csv")
            self.assertIn("configuration_json", settings["key"].tolist())
            self.assertIn("resolution_confirmed", settings["key"].tolist())
            self.assertTrue((Path(directory) / "Unreliable_Spectra.csv").exists())
            book.close()

    def test_row_limit_saves_full_csv_and_never_truncates_workbook(self):
        raw = pd.DataFrame({"signal": np.arange(8, dtype=float)})
        with tempfile.TemporaryDirectory() as directory:
            with patch.object(reports, "EXCEL_MAX_ROWS", 8):
                with self.assertRaisesRegex(ValueError, "not truncated"):
                    reports.export_results(Path(directory), {"Raw_Spectra": raw}, {})
            self.assertEqual(len(pd.read_csv(Path(directory) / "Raw_Spectra.csv")), 8)
            self.assertFalse((Path(directory) / "Amide_I_Results.xlsx").exists())

    def test_oversized_text_fails_explicitly_with_full_csv(self):
        text = "a" * 32768
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError, "text cell"):
                reports.export_results(Path(directory), {
                    "Metadata": pd.DataFrame({"description": [text]})}, {})
            saved = pd.read_csv(Path(directory) / "Metadata.csv")
            self.assertEqual(saved["description"].iloc[0], text)

    def test_settings_preserve_software_metadata(self):
        settings = pd.DataFrame({"key": ["numpy_version", "timestamp_utc"],
                                 "value": [np.__version__, "2026-10-14T00:00:00Z"]})
        with tempfile.TemporaryDirectory() as directory:
            reports.export_results(Path(directory), {"Processing_Settings": settings}, _config())
            values = pd.read_csv(Path(directory) / "Processing_Settings.csv")
            self.assertIn("numpy_version", values["key"].tolist())
            self.assertIn("configuration_json", values["key"].tolist())


class FigureTests(unittest.TestCase):
    def _capture(self):
        figures = {}

        def save(fig, output_dir, stem, config):
            figures[stem] = fig
            return []

        return figures, save

    def tearDown(self):
        plt.close("all")

    def test_group_means_sd_animals_and_adjusted_significance_only(self):
        animals = pd.DataFrame({
            "animal_id": ["c1", "c2", "t1", "t2", "bad"],
            "group": ["Control", "Control", "Cancer", "Cancer", "Cancer"],
            "status": ["Reliable fit"] * 4 + ["Unreliable fit"],
            "random_coil": [10., 20., 30., 40., 99.],
            "alpha_helix": [20., 30., 40., 50., 99.],
            "beta_sheet": [30., 20., 15., 5., 99.],
            "beta_turn": [40., 30., 15., 5., 99.],
        })
        tests = pd.DataFrame({"metric": ["random_coil", "alpha_helix", "beta_sheet"],
                              "group1": ["Control"] * 3, "group2": ["Cancer"] * 3,
                              "p_value": [.0001, .0001, .0001], "q_value": [.08, .001, .0001],
                              "interpretation_allowed": [True, True, False]})
        figures, capture = self._capture()
        with patch.object(plotting, "_save", side_effect=capture):
            plotting.plot_groups(animals, tests, Path("unused"), _config())
        axis = figures["Secondary_Structure_Groups"].axes[0]
        self.assertEqual(len(axis.patches), 8)
        self.assertEqual([patch.get_height() for patch in axis.patches[:2]], [15., 35.])
        self.assertEqual(axis.patches[0].get_facecolor()[-1], 0)
        self.assertEqual(len(axis.get_xticklabels()), 4)
        stars = [text.get_text() for text in axis.texts if "*" in text.get_text()]
        self.assertEqual(stars, ["**"])
        self.assertEqual(sum(len(collection.get_offsets()) for collection in axis.collections
                             if hasattr(collection, "get_offsets") and
                             type(collection).__name__ == "PathCollection"), 16)

    def test_merged_category_is_explicit_and_excluded_animals_make_no_bars(self):
        animals = pd.DataFrame({"group": ["Control"], "status": ["Unreliable fit"],
                                "random_coil": [np.nan], "alpha_helix": [np.nan],
                                "beta_sheet": [np.nan], "beta_turn": [np.nan]})
        config = _config()
        config["plots"]["merge_beta_turn"] = True
        config["plots"]["metric_labels"] = {"beta_associated": "β-sheet"}
        figures, capture = self._capture()
        with patch.object(plotting, "_save", side_effect=capture):
            plotting.plot_groups(animals, pd.DataFrame(), Path("unused"), config)
        axis = figures["Secondary_Structure_Groups"].axes[0]
        self.assertEqual(len(axis.patches), 0)
        self.assertIn("β-sheet + β-turn", axis.get_xticklabels()[-1].get_text())
        self.assertTrue(any("No accepted" in text.get_text() for text in axis.texts))

    def test_unreliable_components_use_shape_even_when_area_is_nan(self):
        x = np.linspace(1600, 1700, 51)
        signal = np.exp(-4 * np.log(2) * ((x - 1655) / 16) ** 2)
        processed = {"x": x, "raw": signal, "baseline": np.zeros_like(x),
                     "corrected": signal, "smoothed": signal,
                     "derivatives": {7: np.gradient(np.gradient(signal, x), x)}}
        fit = {"status": "Unreliable fit", "components": [{
            "center": 1655, "amplitude": 1, "fwhm": 16,
            "structure": "alpha_helix", "area": np.nan, "area_pct": np.nan}],
            "warnings": ["Instrument resolution unconfirmed"]}
        figures, capture = self._capture()
        with patch.object(plotting, "_save", side_effect=capture):
            plotting.plot_spectrum({"x": np.arange(500, 4001), "y": np.ones(3501)},
                                   processed, fit, Path("unused"), "mouse", _config())
        original = figures["mouse_full_spectrum"].axes[0].lines[0]
        self.assertEqual(len(original.get_xdata()), 3501)
        axes = figures["mouse_decomposition"].axes
        component = [line for line in axes[0].lines if line.get_label().startswith("1:")][0]
        self.assertNotIn("%", component.get_label())
        np.testing.assert_allclose(component.get_ydata(), signal)
        self.assertIn("Unreliable fit", axes[0].get_title())
        self.assertTrue(any(line.get_label() == "Gaussian sum" for line in axes[0].lines))

    def test_headless_exports_all_three_formats(self):
        with tempfile.TemporaryDirectory() as directory:
            files = plotting.plot_spectrum({"x": [1000, 2000, 3000], "y": [.1, .2, .1]},
                                           None, None, Path(directory), "mouse", _config())
            self.assertEqual({path.suffix for path in files}, {".png", ".pdf", ".svg"})
            self.assertTrue(all(path.stat().st_size > 0 for path in files))

    def test_synthetic_dataset_label_is_visible_in_exported_figure(self):
        config = _config()
        label = "SYNTHETIC DATA — illustrative verification only"
        config["dataset_label"] = label
        with tempfile.TemporaryDirectory() as directory:
            with plotting._context(config):
                fig, axis = plt.subplots()
                axis.plot([1, 2], [1, 2])
                axis.set_title("Existing title stays visible")
                files = plotting._save(fig, Path(directory), "synthetic", config)
            self.assertTrue(any(text.get_text() == label for text in fig.texts))
            self.assertEqual(axis.get_title(), "Existing title stays visible")
            svg = next(path for path in files if path.suffix == ".svg").read_text()
            self.assertIn(label, svg)


if __name__ == "__main__":
    unittest.main()
