"""A rejected Amide I decomposition must not suppress predefined regional tests."""

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from amide_analysis.config import default_config
from amide_analysis.fitting import STRUCTURES
from amide_analysis.pipeline import run_analysis


class CombinedWorkflowTests(unittest.TestCase):
    def test_all_regions_and_requested_excel_table_survive_rejected_decompositions(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "spectra"
            x = np.arange(450., 4000., 2.)
            for group, amplitudes in (("Control", (.02, .03, .04)),
                                      ("Cancer", (.10, .11, .12))):
                folder = source / group
                folder.mkdir(parents=True)
                for index, amplitude in enumerate(amplitudes):
                    y = .2 + amplitude * np.exp(-.5 * ((x - 1650.) / 10.)**2)
                    pd.DataFrame({"wavenumber": x, "absorbance": y}).to_csv(
                        folder / f"mouse_{index}.csv", index=False)

            def rejected(processed, config, *args, **kwargs):
                nominal = {key: 100. if key == "alpha_helix" else 0. for key in STRUCTURES}
                return {"status": "Unreliable fit", "warnings": ["Unstable model areas"],
                    "fit": np.zeros_like(processed["x"]), "residuals": processed["corrected"],
                    "metrics": {"components": 1}, "numerical_fit_reliable": False,
                    "quantitative_accepted": False, "fit_rejection_codes": ["BOOTSTRAP_AREA_INSTABILITY"],
                    "protocol_blockers": [], "percentages": {key: np.nan for key in STRUCTURES},
                    "diagnostic_percentages": nominal, "nominal_percentages": nominal,
                    "components": [], "candidates": [], "model_comparison": [], "sensitivity": []}

            with patch("amide_analysis.pipeline.analyze_band", side_effect=rejected), \
                 patch("amide_analysis.pipeline.plot_spectrum", return_value=[]), \
                 patch("amide_analysis.pipeline.plot_groups", return_value=[]):
                result = run_analysis(source, root / "results", default_config(), progress=lambda _: None)
            tables = result["tables"]
            self.assertTrue(tables["Secondary_Structure"][list(STRUCTURES)].isna().all().all())
            self.assertTrue(tables["Region_Statistics"]["Tested"].all())
            self.assertEqual(tables["Region_Summary"].shape, (9, 7))
            amide = tables["Region_Statistics"].iloc[0]
            self.assertEqual((amide["Control_N"], amide["Cancer_N"]), (3, 3))
            self.assertAlmostEqual(amide["p_value"], .1)
            self.assertGreater(amide["Effect_size"], 0.)
            self.assertEqual(len(tables["Model_Estimates"]), 6)
            self.assertFalse(tables["Model_Estimates"]["quantitative_accepted"].any())
            workbook = load_workbook(result["excel"])
            self.assertEqual(workbook.active.title, "Analysis_Summary")
            summary = workbook["Region_Summary"]
            self.assertEqual((summary.max_row, summary.max_column), (10, 7))
            self.assertIsInstance(summary["E2"].value, float)
            self.assertIn("<0.0001", summary["E2"].number_format)
            self.assertEqual(workbook["Analysis_Summary"]["C3"].hyperlink.target,
                             "#'Region_Summary'!A1")


if __name__ == "__main__":
    unittest.main()
